import os
import docx
from fpdf import FPDF

class DocumentExporterService:
    """Exports corrected documents into PDF, DOCX, TXT, or Markdown formats."""

    @staticmethod
    def export_document(title: str, text: str, format_type: str, output_path: str) -> str:
        format_type = format_type.lower()
        if format_type == 'pdf':
            return DocumentExporterService._export_pdf(title, text, output_path)
        elif format_type == 'docx':
            return DocumentExporterService._export_docx(title, text, output_path)
        elif format_type in ['md', 'markdown']:
            return DocumentExporterService._export_markdown(title, text, output_path)
        else:
            return DocumentExporterService._export_txt(title, text, output_path)

    @staticmethod
    def _export_txt(title: str, text: str, output_path: str) -> str:
        content = f"PALGRIMIZER CORRECTED DOCUMENT: {title}\n" + "=" * 50 + "\n\n" + text
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(content)
        return output_path

    @staticmethod
    def _export_markdown(title: str, text: str, output_path: str) -> str:
        content = f"# {title}\n\n*Corrected by Palgrimizer AI*\n\n---\n\n{text}"
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(content)
        return output_path

    @staticmethod
    def _export_docx(title: str, text: str, output_path: str) -> str:
        doc = docx.Document()
        doc.add_heading(title, level=0)
        p_sub = doc.add_paragraph()
        run = p_sub.add_run("Corrected with Palgrimizer AI Document Correction Engine")
        run.italic = True

        for para in text.split('\n\n'):
            if para.strip():
                doc.add_paragraph(para.strip())

        doc.save(output_path)
        return output_path

    @staticmethod
    def _export_pdf(title: str, text: str, output_path: str) -> str:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        # Title
        pdf.cell(0, 10, title.encode('latin-1', 'replace').decode('latin-1'), new_x="LMARGIN", new_y="NEXT", align="L")
        pdf.set_font("Helvetica", "I", 10)
        pdf.cell(0, 8, "Palgrimizer AI Enterprise Corrected Document", new_x="LMARGIN", new_y="NEXT", align="L")
        pdf.ln(5)

        pdf.set_font("Helvetica", "", 11)
        for paragraph in text.split('\n\n'):
            if paragraph.strip():
                clean_para = paragraph.encode('latin-1', 'replace').decode('latin-1')
                pdf.multi_cell(0, 6, clean_para)
                pdf.ln(3)

        pdf.output(output_path)
        return output_path

doc_exporter = DocumentExporterService()
